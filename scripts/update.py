import argparse, io, json, hashlib, re, time, zipfile
from pathlib import Path
from urllib.parse import urljoin, urlparse, urlunparse, parse_qsl, urlencode
from datetime import datetime, timezone, timedelta, date

import requests
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry
from bs4 import BeautifulSoup
import pandas as pd
from pypdf import PdfReader
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
CFG = json.loads((ROOT / 'sources.json').read_text(encoding='utf-8'))
PROFILE = json.loads((ROOT / 'docs/data/profile.json').read_text(encoding='utf-8'))
ITEMS_PATH = ROOT / 'docs/data/items.json'
STATUS_PATH = ROOT / 'docs/data/source_status.json'
CHAINS_PATH = ROOT / 'docs/data/chains.json'
TZ = timezone(timedelta(hours=8))
NOW = datetime.now(TZ)
TODAY = NOW.date()
CURRENT_YEAR = TODAY.year

SESSION = requests.Session()
retry = Retry(
    total=2, connect=2, read=2, backoff_factor=.6,
    status_forcelist=[429, 500, 502, 503, 504], allowed_methods=['GET']
)
SESSION.mount('https://', HTTPAdapter(max_retries=retry))
SESSION.mount('http://', HTTPAdapter(max_retries=retry))
HEADERS = {
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) '
                  'AppleWebKit/537.36 Chrome/124 Safari/537.36 TeacherExamRadar/8.0'
}


def now_iso():
    return datetime.now(TZ).isoformat(timespec='seconds')


def clean(s):
    return re.sub(r'\s+', ' ', s or '').strip()


def norm_domain(u):
    return urlparse(u).netloc.lower().split(':')[0].removeprefix('www.')


def same_domain(a, b):
    da, db = norm_domain(a), norm_domain(b)
    return bool(da and db and (da == db or da.endswith('.' + db) or db.endswith('.' + da)))


def canonical_url(u):
    try:
        p = urlparse(u)
        keep = []
        for k, v in parse_qsl(p.query, keep_blank_values=True):
            if not (k.lower().startswith('utm_') or k.lower() in {'spm'}):
                keep.append((k, v))
        path = re.sub(r'/+', '/', p.path or '/')
        return urlunparse((p.scheme.lower(), p.netloc.lower(), path.rstrip('/') or '/', '', urlencode(keep), ''))
    except Exception:
        return u


def get(url, binary=False):
    r = SESSION.get(url, headers=HEADERS, timeout=(10, 30), allow_redirects=True)
    r.raise_for_status()
    if binary:
        return r.content
    r.encoding = r.apparent_encoding or r.encoding
    return r.text


def parse_ymd(y, m, d):
    try:
        return f'{int(y):04d}-{int(m):02d}-{int(d):02d}'
    except Exception:
        return ''


def parse_date(s):
    try:
        return datetime.strptime(s, '%Y-%m-%d').date()
    except Exception:
        return None


def guess_date(title, url, body='', soup=None):
    blobs = []
    if soup:
        for sel in [
            'meta[name="PubDate"]', 'meta[name="publishdate"]', 'meta[name="publishDate"]',
            'meta[property="article:published_time"]', 'time'
        ]:
            try:
                for n in soup.select(sel):
                    blobs.append(n.get('content') or n.get_text(' ', strip=True))
            except Exception:
                pass
    # Prefer an explicitly labelled publication date over arbitrary dates such as registration deadlines in the body.
    m = re.search(r'(?:发布日期|发布时间|发布于|信息时间|成文日期)[：:\s]*(20\d{2})[-年/.](\d{1,2})[-月/.](\d{1,2})日?', (body or '')[:5000])
    if m:
        blobs.append('PUB:'+parse_ymd(*m.groups()))
    blobs.extend([title or '', url or '', (body or '')[:3000]])
    for blob in blobs:
        for p in [
            r'(20\d{2})[-年/.](\d{1,2})[-月/.](\d{1,2})日?',
            r'/(20\d{2})/(\d{1,2})/(\d{1,2})/'
        ]:
            m = re.search(p, blob)
            if m:
                return parse_ymd(*m.groups())
    return ''


def detect_cycle_year(title, full_text='', pub_date=''):
    blob = (title or '') + '\n' + (full_text or '')[:8000]
    m = re.search(r'(20\d{2})\s*届', blob)
    if m:
        return m.group(1)
    m = re.search(r'(20\d{2})\s*年度', title or '')
    if m:
        return m.group(1)
    m = re.search(r'(20\d{2})\s*年', title or '')
    if m:
        return m.group(1)
    return (pub_date[:4] if pub_date else '')


def is_recruitment_candidate(title):
    t = clean(title)
    if len(t) < 5:
        return False
    action = [
        '招聘', '选聘', '引进', '招录', '新任教师', '事业单位', '校园招聘',
        '资格复审', '资格审查', '笔试', '面试', '专业测试', '体检', '考察',
        '拟聘', '拟录聘', '递补', '补录', '报名人数', '报名统计', '核减', '取消招聘计划'
    ]
    if any(k in t for k in action):
        return True
    if '教师' in t and any(k in t for k in ['公告', '报名', '成绩', '名单', '岗位', '考试', '公示']):
        return True
    return False


def detect_stage(title):
    rules = [
        ('递补', ['递补', '补录']),
        ('拟聘', ['拟聘', '拟录聘', '聘用公示', '录用公示']),
        ('体检考察', ['体检', '考察']),
        ('资格复审', ['资格复审', '资格审查']),
        ('面试', ['面试', '专业测试', '试讲', '说课']),
        ('笔试成绩', ['笔试成绩', '成绩查询', '笔试成绩公布', '笔试成绩公示']),
        ('报名统计', ['报名人数', '报名统计', '缴费人数', '报名情况', '开考比例', '核减', '取消招聘计划']),
        ('招聘公告', ['招聘公告', '公开招聘', '校园招聘', '选聘', '引进', '招聘教师', '新任教师招聘'])
    ]
    for stage, keys in rules:
        if any(k in title for k in keys):
            return stage
    return '其他相关'


def detect_recruitment_type(title, text=''):
    blob = clean((title or '') + ' ' + (text or '')[:1200])
    if any(k in blob for k in ['校园招聘', '赴高校招聘', '高校专项招聘', '应届毕业生招聘']):
        return '校招'
    if any(k in blob for k in ['人才引进', '高层次人才', '紧缺人才', '优秀毕业生选聘', '选聘优秀毕业生']):
        return '人才引进/选聘'
    if '统一公开招聘' in blob or '统一招聘' in blob:
        return '统招'
    if '新任教师公开招聘' in blob:
        return '新任教师公开招聘'
    if '面向社会' in blob and '招聘' in blob:
        return '社招/公开招聘'
    if '公开招聘' in blob:
        return '公开招聘'
    return '其他/待核'


def extract_target_numbers(target_rows):
    text = '\n'.join(target_rows or [])
    out = {}
    if not text:
        return out
    compact = re.sub(r'\s+', ' ', text)
    recruit_patterns = [
        r'(?:初中语文|初级中学.{0,8}语文|语文教师).{0,45}?(?:招聘计划|招聘人数|计划招聘|招聘数)?[：:\s|,，]*(\d{1,3})\s*(?:名|人)',
        r'(?:招聘计划|招聘人数|计划招聘|招聘数).{0,25}?(?:初中语文|语文教师).{0,25}?(\d{1,3})\s*(?:名|人)'
    ]
    for pat in recruit_patterns:
        m = re.search(pat, compact)
        if m:
            out['target_subject_recruit'] = int(m.group(1)); break
    cutoff_patterns = [
        r'(?:初中语文|语文教师).{0,80}?(?:最低进面分|最低入围分|最低分|笔试最低分)[：:\s]*(\d+(?:\.\d+)?)',
        r'(?:最低进面分|最低入围分|最低分|笔试最低分)[：:\s]*(\d+(?:\.\d+)?).{0,80}?(?:初中语文|语文教师)'
    ]
    for pat in cutoff_patterns:
        m = re.search(pat, compact)
        if m:
            out['target_cutoff'] = float(m.group(1)); break
    applicants_patterns = [
        r'(?:初中语文|语文教师).{0,80}?(?:报名人数|审核通过人数|缴费人数)[：:\s]*(\d+)',
        r'(?:报名人数|审核通过人数|缴费人数)[：:\s]*(\d+).{0,80}?(?:初中语文|语文教师)'
    ]
    for pat in applicants_patterns:
        m = re.search(pat, compact)
        if m:
            out['target_applicants'] = int(m.group(1)); break
    if out.get('target_subject_recruit') and out.get('target_applicants'):
        out['target_ratio'] = round(out['target_applicants'] / out['target_subject_recruit'], 2)
    return out


DATE_RE = re.compile(
    r'(20\d{2})年(\d{1,2})月(\d{1,2})日|'
    r'(20\d{2})[-/.](\d{1,2})[-/.](\d{1,2})|'
    r'(?<!\d)(\d{1,2})月(\d{1,2})日'
)


def date_matches(blob, fallback_year=None):
    vals = []
    inherited = int(fallback_year) if str(fallback_year).isdigit() else None
    for m in DATE_RE.finditer(blob or ''):
        g = m.groups()
        if g[0]:
            inherited = int(g[0]); vals.append(parse_ymd(g[0], g[1], g[2]))
        elif g[3]:
            inherited = int(g[3]); vals.append(parse_ymd(g[3], g[4], g[5]))
        elif inherited:
            vals.append(parse_ymd(inherited, g[6], g[7]))
    vals = [x for x in vals if x]
    # Common shorthand: “10月8日9:00至12日16:00”. Infer the second day in the same month/year.
    if len(vals) == 1:
        m = re.search(r'(\d{1,2})月(\d{1,2})日.{0,28}?(?:至|到|—|–|-|~|～)(\d{1,2})日', blob or '')
        if m and inherited:
            second = parse_ymd(inherited, m.group(1), m.group(3))
            if second and second != vals[0]:
                vals.append(second)
    return vals


def extract_schedule(text, fallback_year=None):
    out, t = {}, text or ''

    def first_after(patterns, chars=260):
        for pat in patterns:
            m = re.search(pat, t)
            if not m:
                continue
            frag = t[m.start():m.start() + chars]
            vals = date_matches(frag, fallback_year)
            if vals:
                return vals, frag
        return [], ''

    vals, frag = first_after([r'报名时间', r'网上报名(?:及材料上传)?时间', r'报名及资格初审', r'报名采用', r'报名方式及时间'])
    if vals:
        out['apply_start'] = vals[0]
        if len(vals) > 1:
            out['apply_end'] = vals[1]
        elif '截止' in frag:
            out['apply_end'] = vals[0]

    vals2, _ = first_after([r'报名截止(?:时间)?', r'网上报名截止(?:时间)?', r'截止报名'])
    if vals2:
        out['apply_end'] = vals2[0]

    vals, _ = first_after([r'笔试时间', r'笔试日期', r'统一笔试'])
    if vals:
        out['written_start'] = vals[0]

    vals, _ = first_after([r'面试时间', r'专业测试时间', r'试讲时间', r'说课时间'])
    if vals:
        out['interview_start'] = vals[0]

    return out


def application_status(schedule, pub_date='', stage=''):
    start = parse_date((schedule or {}).get('apply_start', ''))
    end = parse_date((schedule or {}).get('apply_end', ''))
    if end and end < TODAY:
        return '已截止'
    if start and start > TODAY:
        return '即将报名'
    if start and (not end or end >= TODAY) and start <= TODAY:
        return '报名中'
    if end and end >= TODAY:
        return '报名中/时间待核'
    if stage == '招聘公告':
        d = parse_date(pub_date)
        if d and (TODAY - d).days <= 90:
            return '近期公告·报名时间待核'
    return '时间待核'


def classify(stage, pub_date, cycle_year, schedule):
    status = application_status(schedule, pub_date, stage)
    if stage in {'拟聘', '体检考察', '资格复审', '面试', '笔试成绩', '递补'}:
        return '历史考情'
    if stage == '招聘公告':
        if status in {'报名中', '即将报名', '报名中/时间待核', '近期公告·报名时间待核'}:
            return '近期公告'
        # If the site hides/omits the publication date and schedule, do not bury a current/next-cohort notice in history.
        if status == '时间待核' and not pub_date and str(cycle_year) in {str(CURRENT_YEAR), str(CURRENT_YEAR + 1)}:
            return '近期公告'
        return '历史考情'
    if stage == '报名统计':
        d = parse_date(pub_date)
        if d and (TODAY - d).days <= 45:
            return '近期公告'
    return '历史考情'


def extract_numbers(text):
    t = (text or '').replace(',', '').replace('，', '')
    out = {}
    rules = {
        'total_recruit': [
            r'计划(?:面向社会)?(?:公开)?招聘(?:教师|工作人员)?\s*(\d+)\s*名',
            r'招聘计划(?:岗位)?(?:共|为)?\s*(\d+)\s*名',
            r'公开招聘(?:教师)?\s*(\d+)\s*名'
        ],
        'applicants': [r'报名(?:总)?人数[：:\s]*(\d+)', r'共有\s*(\d+)\s*人报名', r'报名成功\s*(\d+)\s*人', r'报名人数为\s*(\d+)'],
        'qualified': [r'(?:资格)?审核通过(?:人数)?[：:\s]*(\d+)', r'通过资格(?:初审|审核)\s*(\d+)\s*人'],
        'attended': [r'实际(?:参加)?(?:笔试|考试)(?:人数)?[：:\s]*(\d+)', r'参考人数[：:\s]*(\d+)'],
        'interview_count': [r'进入(?:专业测试|面试)(?:人数)?[：:\s]*(\d+)', r'入围(?:专业测试|面试)\s*(\d+)\s*人'],
        'admitted': [r'拟聘用(?:人员)?(?:共)?\s*(\d+)\s*名', r'拟录聘(?:人员)?(?:共)?\s*(\d+)\s*名']
    }
    for key, pats in rules.items():
        for p in pats:
            m = re.search(p, t)
            if m:
                out[key] = int(m.group(1)); break
    if out.get('applicants') and out.get('total_recruit'):
        out['overall_ratio'] = round(out['applicants'] / out['total_recruit'], 2)
    for p in [r'最低(?:进面|入围)(?:分|分数|分数线|成绩)?[：:\s]*(\d+(?:\.\d+)?)', r'最低笔试(?:成绩|分数)[：:\s]*(\d+(?:\.\d+)?)']:
        m = re.search(p, t)
        if m:
            out['cutoff'] = float(m.group(1)); break
    return out


def text_lines(text):
    return [clean(x) for x in re.split(r'[\n\r]+', text or '') if clean(x)]


def extract_target_rows(text):
    lines = text_lines(text)
    rows = []
    for i, line in enumerate(lines):
        compact = re.sub(r'\s+', '', line)
        if '语文' not in compact:
            continue
        if not any(x in compact for x in ['初中', '初级中学', '中学', '教师', '学科', '岗位', '专业', '研究生', '硕士', '不限', '教育']):
            continue
        context = [line]
        # A spreadsheet/HTML-table row already carries its own cells. Pulling adjacent rows would mix
        # requirements from another subject (for example, mathematics) into the Chinese job.
        tabular = ('\t' in line or line.count('|') >= 2 or line.count(',') >= 3)
        if not tabular:
            for j in range(max(0, i - 1), min(len(lines), i + 3)):
                if j == i:
                    continue
                c = lines[j]
                if any(k in c for k in ['专业', '学历', '学位', '资格证', '教师资格', '本科', '研究生', '硕士', '岗位代码', '学科']):
                    context.append(c)
        rows.append(' ｜ '.join(context)[:900])
    seen, out = set(), []
    for x in rows:
        if x not in seen:
            seen.add(x); out.append(x)
    out.sort(key=lambda x: (0 if '初中语文' in re.sub(r'\s+', '', x) else 1, len(x)))
    return out[:20]


def extract_cert_rows(text):
    out = []
    for line in text_lines(text):
        if '教师资格' in line and any(x in line for x in ['语文', '初中', '初级中学', '高中', '相应学科', '相应教师资格']):
            out.append(line[:700])
    return list(dict.fromkeys(out))[:10]


def cert_match(cert_rows, target_rows):
    t = '\n'.join(cert_rows + target_rows)
    compact = re.sub(r'\s+', '', t)
    if any(x in compact for x in ['学科学段与招聘岗位一致', '学段学科与招聘岗位一致', '学段、学科与招聘岗位一致', '学段和学科与招聘岗位一致']):
        return '教资需核：公告强调学段/学科与岗位一致；已有高中语文教资，报初中仍需按当地资格审查口径确认'
    if '语文' in compact and any(x in compact for x in ['初中及以上', '初级中学及以上']):
        return '教资匹配：已有高中语文教资，满足“初中及以上语文”层级表述'
    if '高中语文教师资格' in compact:
        return '教资匹配：已有高中语文教资'
    if '相应教师资格' in compact or '相应学科教师资格' in compact:
        return '教资需核：已有高中语文教资，但公告写“相应”时仍按岗位原文确认'
    return '教资要求待核'


def match_profile(full_text, target_rows):
    focus = '\n'.join(target_rows) if target_rows else ''
    t = focus or (full_text or '')
    compact = re.sub(r'\s+', '', t)
    bach = PROFILE.get('bachelor_major', '小学教育')
    mast = PROFILE.get('master_major', '现代教育技术')
    if focus:
        has_wide_grad = ('不限专业' in compact or '专业不限' in compact) and ('研究生' in compact or '硕士' in compact)
        has_bachelor_restrict = any(x in compact for x in [
            '本科所学专业须为与招聘岗位相对应', '本科阶段所学专业须为与招聘岗位相对应',
            '本科阶段所学专业为与招聘岗位相对应', '本科专业与招聘岗位相对应',
            '本科阶段专业与招聘岗位相对应', '本科所学专业与招聘岗位一致', '本科专业须与岗位一致'
        ])
        has_chinese_restrict = any(x in compact for x in [
            '汉语言文学', '中国语言文学', '学科教学（语文）', '学科教学(语文)',
            '课程与教学论（语文）', '课程与教学论(语文)'
        ])
        # Prefer caution when broad and restrictive clauses coexist.
        if has_bachelor_restrict and has_wide_grad:
            return '需重点核实：同一语文岗位同时出现研究生专业不限与本科阶段专业对应要求'
        if has_bachelor_restrict:
            return '专业可能受限：语文岗位要求本科阶段专业与招聘岗位对应'
        if has_chinese_restrict and not has_wide_grad:
            return '专业可能受限：语文岗位出现中文/语文专业限定'
        if has_wide_grad:
            return '高度关注：语文相关行出现研究生/硕士专业不限'
        if any(x in compact for x in ['本科或研究生', '本科或硕士', '任一阶段', '其中一个学历阶段', '本科阶段或研究生阶段']):
            return '高度关注：语文相关行可能允许本科/硕士任一阶段满足'
        if bach in compact or mast in compact:
            return '较匹配：语文相关行出现你的专业名称'
        if any(x in compact for x in ['教育学类', '教育类']) and '语文' in compact:
            return '值得核实：语文相关行出现教育类专业'
        if has_chinese_restrict:
            return '专业可能受限：语文岗位出现中文/语文专业限定'
        return '值得核实：已找到语文岗位相关行'
    if '2027届' in compact and ('教师' in compact or '教育' in compact):
        return '值得核实：面向2027届，但需定位初中语文岗位'
    return '需要核实岗位表'

def extract_exam_clues(text):
    keys = ['教育综合知识', '学科专业知识', '综合应用能力', '职业能力倾向测验', '试讲', '说课', '答辩', '结构化面试', '模拟上课', '笔试成绩', '面试成绩']
    chunks = [clean(x) for x in re.split(r'[。；;\n\r]+', text or '') if clean(x)]
    out = []
    for c in chunks:
        if any(k in c for k in keys) or re.search(r'(?:笔试|面试|专业测试).{0,18}\d{1,3}%|\d{1,3}%.{0,18}(?:笔试|面试|专业测试)', c):
            if len(c) <= 500:
                out.append(c)
    return list(dict.fromkeys(out))[:8]


def is_attachment_link(url, name=''):
    blob = ((url or '') + ' ' + (name or '')).lower()
    if any(ext in blob for ext in CFG['attachment_exts']):
        return True
    if any(k in blob for k in ['download', 'downfile', 'attachment']) and any(k in (name or '') for k in ['附件', '岗位', '名单', '计划表', '信息表', '成绩表']):
        return True
    return False


def attachment_ext(url, name=''):
    low = ((url or '') + ' ' + (name or '')).lower()
    # Longer extensions first so .xlsx is not mistaken for .xls, etc.
    for ext in ['.xlsx', '.docx', '.pdf', '.xls', '.csv', '.doc', '.wps', '.zip', '.rar']:
        if ext in low:
            return ext
    return ''


def infer_attachment_ext(url, name, headers, data):
    ext = attachment_ext(url, name)
    if ext:
        return ext
    cd = (headers.get('Content-Disposition') or headers.get('content-disposition') or '')
    ct = (headers.get('Content-Type') or headers.get('content-type') or '').lower()
    ext = attachment_ext('', cd)
    if ext:
        return ext
    if data.startswith(b'%PDF'):
        return '.pdf'
    if data.startswith(b'PK'):
        if 'spreadsheet' in ct or 'excel' in ct:
            return '.xlsx'
        if 'wordprocessingml' in ct or 'word' in ct:
            return '.docx'
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                names = zf.namelist()
                if any(n.startswith('xl/') for n in names):
                    return '.xlsx'
                if any(n.startswith('word/') for n in names):
                    return '.docx'
        except Exception:
            pass
    if data.startswith(bytes.fromhex('D0CF11E0A1B11AE1')):
        return '.xls' if ('excel' in ct or 'spreadsheet' in ct) else '.doc'
    if 'text/csv' in ct or 'csv' in ct:
        return '.csv'
    return ''


def download_with_meta(url, max_bytes=30 * 1024 * 1024):
    r = SESSION.get(url, headers=HEADERS, timeout=(10, 35), allow_redirects=True, stream=True)
    r.raise_for_status()
    headers = dict(r.headers)
    try:
        clen = int(headers.get('Content-Length') or headers.get('content-length') or 0)
    except Exception:
        clen = 0
    if clen and clen > max_bytes:
        raise ValueError(f'附件过大，跳过自动解析（{clen} bytes）')
    chunks, total = [], 0
    for chunk in r.iter_content(chunk_size=1024 * 256):
        if not chunk:
            continue
        total += len(chunk)
        if total > max_bytes:
            raise ValueError(f'附件超过自动解析上限（{max_bytes} bytes）')
        chunks.append(chunk)
    return b''.join(chunks), headers


def parse_attachment(url, name=''):
    hinted = attachment_ext(url, name)
    if hinted in {'.doc', '.wps', '.zip', '.rar'}:
        return ''
    data, headers = download_with_meta(url)
    ext = infer_attachment_ext(url, name, headers, data)
    if ext not in {'.pdf', '.xlsx', '.xls', '.csv', '.docx'}:
        return ''
    try:
        if ext == '.pdf':
            r = PdfReader(io.BytesIO(data))
            return '\n'.join((p.extract_text() or '') for p in r.pages[:100])
        if ext in {'.xlsx', '.xls'}:
            xls = pd.ExcelFile(io.BytesIO(data)); chunks = []
            for sheet in xls.sheet_names[:12]:
                df = pd.read_excel(io.BytesIO(data), sheet_name=sheet, header=None, dtype=str).fillna('').astype(str)
                # Keep each worksheet row on one logical line so a job title stays attached to its major/degree cells.
                df = df.map(lambda v: clean(str(v).replace('\r', ' ').replace('\n', ' ')))
                rows = ['\t'.join(row) for row in df.values.tolist()]
                chunks.append('\n'.join(rows))
            return '\n'.join(chunks)
        if ext == '.csv':
            try:
                txt = data.decode('utf-8-sig')
            except UnicodeDecodeError:
                txt = data.decode('gb18030', 'ignore')
            return '\n'.join(clean(line) for line in txt.splitlines())
        if ext == '.docx':
            d = Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs]
            parts += ['\t'.join(clean(c.text) for c in row.cells) for table in d.tables for row in table.rows]
            return '\n'.join(parts)
    except Exception:
        return ''
    return ''

def extract_registration_links(base_url, soup, body_raw=''):
    out = []
    keys = ['报名入口', '报名网址', '报名网站', '网上报名', '报名系统', '考试服务平台', '人事考试平台', '人事考试网']
    for a in soup.find_all('a'):
        text = clean(a.get_text(' ', strip=True))
        href = element_link(base_url, a)
        if not href:
            continue
        if not href.startswith(('http://', 'https://')) or is_attachment_link(href, text):
            continue
        blob = text + ' ' + href.lower()
        if any(k in text for k in keys) or any(k in href.lower() for k in ['/signup', '/register', '/apply', 'baoming']):
            out.append({'name': text or '报名入口', 'url': href})
    # Some government pages print a plain URL after “报名网址：” without an anchor.
    for m in re.finditer(r'(?:报名网址|报名网站|报名入口)[：:\s]*((?:https?://|www\.)[^\s，。；;<>]+)', body_raw or '', re.I):
        href = m.group(1).rstrip('）)]}')
        if href.startswith('www.'):
            href = 'https://' + href
        out.append({'name': '报名入口', 'url': canonical_url(href)})
    seen, uniq = set(), []
    for x in out:
        if x['url'] == canonical_url(base_url) or x['url'] in seen:
            continue
        seen.add(x['url']); uniq.append(x)
    return uniq[:8]


def element_link(base_url, a):
    raw = a.get('href') or a.get('data-href') or a.get('data-url') or a.get('data-download') or ''
    if (not raw or raw.lower().startswith('javascript')) and a.get('onclick'):
        m = re.search(r"[\"'](https?://[^\"']+|/[^\"']+)[\"']", a.get('onclick') or '')
        if m:
            raw = m.group(1)
    return canonical_url(urljoin(base_url, raw)) if raw else ''


def article(url, title=''):
    html = get(url)
    soup = BeautifulSoup(html, 'html.parser')
    for tag in soup.find_all(['script', 'style', 'noscript', 'svg']):
        tag.decompose()
    body_raw = soup.get_text('\n', strip=True)
    body = clean(body_raw)
    # Preserve HTML table rows as single logical lines; otherwise BeautifulSoup emits one cell per line
    # and a Chinese-job cell can become detached from its major/degree requirement cells.
    table_rows = []
    for tr in soup.find_all('tr'):
        cells = [clean(c.get_text(' ', strip=True)) for c in tr.find_all(['th', 'td'])]
        cells = [c for c in cells if c]
        if cells:
            table_rows.append(' | '.join(cells))
    merged = body_raw + ('\n' + '\n'.join(table_rows) if table_rows else '')
    registration_links = extract_registration_links(url, soup, body_raw)
    attachments = []
    for a in soup.find_all('a'):
        href = element_link(url, a)
        if not href or not href.startswith(('http://', 'https://')):
            continue
        name = clean(a.get_text(' ', strip=True)) or href.split('/')[-1]
        if is_attachment_link(href, name):
            txt, err = '', ''
            try:
                txt = parse_attachment(href, name)
            except Exception as e:
                err = str(e)
            attachments.append({'name': name, 'url': href, 'parsed': bool(txt), 'preview': clean(txt[:700]), 'error': err})
            if txt:
                merged += '\n' + txt[:180000]
    pub = guess_date(title, url, body_raw, soup)
    return body, merged, attachments, pub, registration_links


def normalize_chain_title(title):
    t = title or ''
    for x in [
        '拟聘用人员名单', '拟聘用人员公示', '拟录聘人员名单', '体检公告', '体检考察', '资格复审',
        '笔试成绩', '成绩查询', '面试公告', '专业测试', '报名情况', '报名人数', '递补', '核减', '取消招聘计划',
        '公告', '通知', '公示', '人员名单', '结果', '关于'
    ]:
        t = t.replace(x, '')
    t = re.sub(r'[（(].*?[）)]', '', t)
    t = re.sub(r'20\d{2}(?:年|年度|届)?', '', t)
    t = t.replace('年度', '').replace('届', '')
    t = clean(t).strip(' -—：:·')[:52]
    return t or '教师招聘'


def chain_key(region, cycle_year, title):
    return f'{region}|{cycle_year or "未知年份"}|{normalize_chain_title(title)}'


def is_pagination(text, href):
    tx = clean(text)
    return tx in ['下一页', '下页', '后一页', '尾页', '上一页', '前一页'] or bool(
        re.search(r'(?:index[_-]?\d+\.html|[?&](?:page|p|pageNo|current)=\d+)', href, re.I)
    )


def nav_candidates(base_url, soup):
    scored = []
    high_words = ['教师招考', '教师招聘', '公开招聘', '校园招聘', '招考', '招聘', '招录', '人事信息']
    for a in soup.find_all('a', href=True):
        text = clean(a.get_text(' ', strip=True))
        href = canonical_url(urljoin(base_url, a['href']))
        if not same_domain(base_url, href):
            continue
        if is_attachment_link(href, text):
            continue
        short_nav = (len(text) <= 14 and any(k == text or k in text for k in CFG['nav_keywords']))
        pu = urlparse(href)
        nav_blob = ((pu.path or '') + '?' + (pu.query or '')).lower()
        # Do not inspect the hostname here: domains such as jytyj.* would otherwise make every link a nav link.
        url_nav = any(x in nav_blob for x in ['/col/', '/class/', '/list', 'list?', 'tzgg', 'gsgg', 'xxgk', 'renshi', 'zhaopin', 'zpzk'])
        score = 0
        if any(w in text for w in high_words):
            score = 120
        elif short_nav:
            score = 90
        if is_pagination(text, href):
            score = max(score, 80)
        elif url_nav:
            score = max(score, 55)
        if score:
            scored.append((score, href))
    # Stable de-duplication after relevance sorting, so limited deep crawls reach recruitment columns first.
    out, seen = [], set()
    for _, href in sorted(scored, key=lambda x: -x[0]):
        if href not in seen:
            seen.add(href); out.append(href)
    return out


def infer_schedule_year(pub_date, cycle_year, title='', full_text=''):
    """Choose the calendar year for month/day-only schedule strings.

    Next-cohort campus recruitment is commonly published in the previous calendar
    year (for example, 2027届 in Oct 2026). If publication metadata is absent,
    prefer the current crawl year for that pattern instead of shifting dates forward.
    """
    if pub_date and len(pub_date) >= 4 and pub_date[:4].isdigit():
        return pub_date[:4]
    cy = int(cycle_year) if str(cycle_year).isdigit() else None
    blob = (title or '') + '\n' + (full_text or '')[:3000]
    if cy == CURRENT_YEAR + 1 and (f'{cy}届' in blob or '校园招聘' in blob or '应届毕业生' in blob):
        return str(CURRENT_YEAR)
    if cy:
        return str(cy)
    return str(CURRENT_YEAR)


def build_item(src, title, href, old_item=None):
    href = canonical_url(href)
    try:
        body, merged, at, pub, reg_links = article(href, title)
        cycle_year = detect_cycle_year(title, merged, pub)
        fallback_year = infer_schedule_year(pub, cycle_year, title, merged)
        stage = detect_stage(title)
        schedule = extract_schedule(merged, fallback_year)
        ex = extract_numbers(merged)
        rows = extract_target_rows(merged)
        ex.update(extract_target_numbers(rows))
        cert_rows = extract_cert_rows(merged)
        iid = (old_item or {}).get('id') or hashlib.sha1((href + '|' + title).encode('utf-8')).hexdigest()[:20]
        item = {
            'id': iid, 'title': title, 'region': src['region'], 'source': src['name'], 'source_id': src['id'],
            'url': href, 'date': pub, 'publish_year': pub[:4] if pub else '', 'cycle_year': cycle_year,
            'year': cycle_year, 'stage': stage, 'recruitment_type': detect_recruitment_type(title, body), 'official': True, 'summary': body[:650], 'attachments': at, 'registration_links': reg_links,
            'extracted': ex, 'schedule': schedule, 'application_status': application_status(schedule, pub, stage),
            'target_rows': rows, 'cert_rows': cert_rows, 'cert_match': cert_match(cert_rows, rows),
            'exam_clues': extract_exam_clues(merged), 'machine_match': match_profile(merged, rows),
            'is_2027': (cycle_year == '2027' or '2027届' in merged or '2027届' in title),
            'chain_key': chain_key(src['region'], cycle_year, title), 'last_checked_at': now_iso()
        }
        item['kind'] = classify(stage, pub, cycle_year, schedule)
        item['discovered_at'] = (old_item or {}).get('discovered_at') or now_iso()
        item['content_hash'] = hashlib.sha1((title + '\n' + body + '\n' + json.dumps(at, ensure_ascii=False, sort_keys=True)).encode('utf-8', 'ignore')).hexdigest()
        if old_item and old_item.get('content_hash') and old_item.get('content_hash') != item['content_hash']:
            item['content_changed_at'] = now_iso()
        else:
            item['content_changed_at'] = (old_item or {}).get('content_changed_at', '')
        return item
    except Exception as e:
        if old_item:
            x = dict(old_item); x['last_checked_at'] = now_iso(); x['refresh_error'] = str(e)[:220]; return x
        cycle_year = detect_cycle_year(title, '', guess_date(title, href))
        iid = hashlib.sha1((href + '|' + title).encode('utf-8')).hexdigest()[:20]
        return {
            'id': iid, 'title': title, 'region': src['region'], 'source': src['name'], 'source_id': src['id'],
            'url': href, 'date': guess_date(title, href), 'publish_year': '', 'cycle_year': cycle_year, 'year': cycle_year,
            'stage': detect_stage(title), 'recruitment_type': detect_recruitment_type(title, ''), 'kind': '近期公告' if cycle_year in {str(CURRENT_YEAR), str(CURRENT_YEAR + 1)} else '历史考情',
            'official': True, 'summary': '发现官方链接，但详情抓取失败；请直接打开原公告核验。', 'attachments': [], 'registration_links': [],
            'extracted': {}, 'schedule': {}, 'application_status': '时间待核', 'target_rows': [], 'cert_rows': [],
            'cert_match': '教资要求待核', 'exam_clues': [], 'machine_match': '需要人工核实',
            'is_2027': cycle_year == '2027', 'chain_key': chain_key(src['region'], cycle_year, title),
            'discovered_at': now_iso(), 'last_checked_at': now_iso(), 'fetch_error': str(e)[:220]
        }


def hours_since(iso):
    if not iso:
        return 10 ** 9
    try:
        return (datetime.now(TZ) - datetime.fromisoformat(iso)).total_seconds() / 3600
    except Exception:
        return 10 ** 9


def scan_source(src, items_by_url, prev, mode='quick'):
    checked, candidate_links, fingerprints, page_errors = [], [], [], []
    entries = [canonical_url(x) for x in src.get('entry_urls', [])]
    # Hourly quick scans always start from current entry pages. Daily deep scans resume the saved frontier first,
    # allowing 3–5 year history to progress beyond the first N pages over successive days.
    prior_frontier = [canonical_url(x) for x in (prev.get('pending_pages', []) if mode == 'deep' else [])]
    queue = list(dict.fromkeys(prior_frontier + entries if mode == 'deep' else entries))
    seen_pages = set()
    configured = int(src.get('max_pages', 12))
    max_pages = configured if mode == 'deep' else min(configured, int(src.get('quick_pages', 4)))
    # Keep each job bounded. Deep runs backfill history incrementally instead of trying to fetch years of articles at once.
    new_budget = 25 if mode == 'deep' else 8
    refresh_budget = 15 if mode == 'deep' else 4
    refresh_hours = 168 if mode == 'deep' else 24

    while queue and len(seen_pages) < max_pages:
        page = queue.pop(0)
        if page in seen_pages:
            continue
        seen_pages.add(page)
        try:
            html = get(page); soup = BeautifulSoup(html, 'html.parser'); checked.append(page)
        except Exception as e:
            page_errors.append({'url': page, 'error': str(e)[:180]}); continue

        relevant_snapshot = []
        for a in soup.find_all('a', href=True):
            title = clean(a.get_text(' ', strip=True))
            href = canonical_url(urljoin(page, a['href']))
            if not href.startswith(('http://', 'https://')) or not same_domain(page, href):
                continue
            if title and (any(k in title for k in CFG['discover_keywords']) or is_pagination(title, href)):
                relevant_snapshot.append(title + '|' + href)
            if not is_attachment_link(href, title) and is_recruitment_candidate(title):
                candidate_links.append((title, href))
        fingerprints.append(page + '\n' + '\n'.join(sorted(set(relevant_snapshot))[:900]))
        for nav in nav_candidates(page, soup):
            if nav not in seen_pages and nav not in queue and len(queue) + len(seen_pages) < max_pages * 4:
                queue.append(nav)

    uniq, seen = [], set()
    for t, u in candidate_links:
        k = canonical_url(u)
        if k not in seen:
            seen.add(k); uniq.append((t, k))

    fp = hashlib.sha1('\n'.join(fingerprints).encode('utf-8', 'ignore')).hexdigest()
    prev_fp = prev.get(f'{mode}_fingerprint') or (prev.get('fingerprint') if not prev.get('quick_fingerprint') and not prev.get('deep_fingerprint') else '')
    source_changed = bool(prev_fp and fp != prev_fp)
    updates, new_count, refresh_count, changed_count, detail_errors = [], 0, 0, 0, []
    unseen_total = sum(1 for _, href in uniq if items_by_url.get(href) is None)

    # Iterate the whole discovered list so that yesterday's newly-saved first page does not permanently block older unseen items.
    for title, href in uniq:
        old = items_by_url.get(href)
        if old is None:
            if new_count >= new_budget:
                continue
            item = build_item(src, title, href); new_count += 1; updates.append(item)
            if item.get('fetch_error'):
                detail_errors.append({'url': href, 'title': title, 'error': item.get('fetch_error')})
        elif (source_changed or hours_since(old.get('last_checked_at')) >= refresh_hours) and refresh_count < refresh_budget:
            before = old.get('content_hash')
            item = build_item(src, title, href, old); refresh_count += 1; updates.append(item)
            if item.get('refresh_error'):
                detail_errors.append({'url': href, 'title': title, 'error': item.get('refresh_error')})
            if before and item.get('content_hash') and before != item.get('content_hash'):
                changed_count += 1
        time.sleep(.04)

    return updates, {
        'checked_pages': checked, 'candidate_count': len(uniq), 'fingerprint': fp, 'page_errors': page_errors, 'detail_errors': detail_errors[:20],
        'changed': source_changed, 'new_count': new_count, 'refresh_count': refresh_count, 'changed_count': changed_count,
        'backfill_remaining': max(0, unseen_total - new_count),
        'pending_pages': list(dict.fromkeys(queue))[:100] if mode == 'deep' else prev.get('pending_pages', []),
        'pagination_pending': len(queue) if mode == 'deep' else len(prev.get('pending_pages', [])),
        'mode': mode
    }


def build_chains(items):
    groups = {}
    for it in items:
        k = it.get('chain_key') or chain_key(it.get('region', ''), it.get('cycle_year') or it.get('year', ''), it.get('title', ''))
        groups.setdefault(k, []).append(it)
    chains = []
    order = ['招聘公告', '报名统计', '笔试成绩', '资格复审', '面试', '体检考察', '拟聘', '递补', '其他相关']
    for k, arr in groups.items():
        arr = sorted(arr, key=lambda x: (x.get('date', ''), order.index(x.get('stage')) if x.get('stage') in order else 99))
        stages = []
        for x in arr:
            if x.get('stage') not in stages:
                stages.append(x.get('stage'))
        nums = {}
        for x in arr:
            for nk, nv in (x.get('extracted') or {}).items():
                if nv is not None:
                    nums[nk] = nv
        root = next((x for x in arr if x.get('stage') == '招聘公告'), arr[0])
        chains.append({
            'chain_key': k, 'region': root.get('region', ''), 'year': root.get('cycle_year') or root.get('year', ''),
            'title': root.get('title', ''), 'recruitment_type': root.get('recruitment_type', ''), 'stages': stages, 'item_count': len(arr),
            'latest_date': max((x.get('date', '') for x in arr), default=''), 'numbers': nums,
            'items': [{'title': x.get('title'), 'stage': x.get('stage'), 'date': x.get('date'), 'url': x.get('url')} for x in arr]
        })
    return sorted(chains, key=lambda x: (x.get('latest_date', ''), x.get('region', '')), reverse=True)


def main(mode='quick'):
    old = {'items': []}
    if ITEMS_PATH.exists():
        try:
            old = json.loads(ITEMS_PATH.read_text(encoding='utf-8'))
        except Exception:
            pass
    prev_status = {}
    if STATUS_PATH.exists():
        try:
            prev_status = {x['source_id']: x for x in json.loads(STATUS_PATH.read_text(encoding='utf-8')).get('sources', [])}
        except Exception:
            pass

    items = old.get('items', [])
    by_url = {canonical_url(x.get('url')): x for x in items if x.get('url')}
    statuses = []

    for src in CFG['sources']:
        p = prev_status.get(src['id'], {})
        started = now_iso()
        perf_started = time.monotonic()
        try:
            updates, meta = scan_source(src, by_url, p, mode)
            for item in updates:
                by_url[canonical_url(item['url'])] = item
            checked = len(meta['checked_pages']); pe = meta['page_errors']
            if checked == 0:
                stat = '抓取失败'
            elif pe or meta.get('detail_errors'):
                stat = '部分失败'
            elif meta['candidate_count'] == 0:
                # A 200 response can still be a JS shell / anti-bot page. Do not call that fully healthy.
                stat = '可访问但未发现候选'
            else:
                stat = '正常'
            changed = meta['changed'] or meta['new_count'] > 0 or meta['changed_count'] > 0
            duration = round(time.monotonic() - perf_started, 2)
            success_now = checked > 0
            consecutive_failures = 0 if success_now else int(p.get('consecutive_failures', 0) or 0) + 1
            last_success_at = started if success_now else p.get('last_success_at', '')
            statuses.append({
                'source_id': src['id'], 'name': src['name'], 'region': src['region'], 'category': src.get('category', ''),
                'priority': src.get('priority', ''), 'status': stat, 'checked_at': started, 'last_success_at': last_success_at,
                'consecutive_failures': consecutive_failures, 'duration_seconds': duration, 'checked_pages': checked,
                'candidate_count': meta['candidate_count'], 'new_items': meta['new_count'], 'refreshed_items': meta['refresh_count'],
                'changed_items': meta['changed_count'],
                'backfill_remaining': meta.get('backfill_remaining', 0) if mode == 'deep' else max(meta.get('backfill_remaining', 0), p.get('backfill_remaining', 0)),
                'pending_pages': meta.get('pending_pages', p.get('pending_pages', [])), 'pagination_pending': meta.get('pagination_pending', p.get('pagination_pending', 0)),
                'changed': changed, 'last_change': started if changed else p.get('last_change', ''),
                'fingerprint': meta['fingerprint'],
                'quick_fingerprint': meta['fingerprint'] if mode == 'quick' else p.get('quick_fingerprint', ''),
                'deep_fingerprint': meta['fingerprint'] if mode == 'deep' else p.get('deep_fingerprint', ''),
                'error': ('；'.join(x['error'] for x in pe[:2]) if checked == 0 and pe else ''),
                'page_errors': pe[:10], 'detail_errors': meta.get('detail_errors', [])[:10], 'entry_urls': src.get('entry_urls', []), 'mode': mode
            })
        except Exception as e:
            statuses.append({
                'source_id': src['id'], 'name': src['name'], 'region': src['region'], 'category': src.get('category', ''),
                'priority': src.get('priority', ''), 'status': '抓取失败', 'checked_at': started, 'last_success_at': p.get('last_success_at', ''),
                'consecutive_failures': int(p.get('consecutive_failures', 0) or 0) + 1, 'duration_seconds': round(time.monotonic() - perf_started, 2), 'checked_pages': 0,
                'candidate_count': 0, 'new_items': 0, 'refreshed_items': 0, 'changed_items': 0, 'backfill_remaining': p.get('backfill_remaining', 0),
                'pending_pages': p.get('pending_pages', []), 'pagination_pending': p.get('pagination_pending', 0), 'changed': False,
                'last_change': p.get('last_change', ''), 'fingerprint': p.get('fingerprint', ''),
                'quick_fingerprint': p.get('quick_fingerprint', ''), 'deep_fingerprint': p.get('deep_fingerprint', ''), 'error': str(e)[:300],
                'page_errors': [], 'detail_errors': [], 'entry_urls': src.get('entry_urls', []), 'mode': mode
            })

    items = list(by_url.values())
    items += [x for x in old.get('items', []) if not x.get('url')]
    merged = {}
    for it in items:
        key = it.get('id') or hashlib.sha1((it.get('url', '') + '|' + it.get('title', '')).encode()).hexdigest()[:20]
        it['id'] = key; merged[key] = it
    items = sorted(merged.values(), key=lambda x: (x.get('date', ''), x.get('discovered_at', ''), x.get('title', '')), reverse=True)[:5000]
    now = now_iso()
    ITEMS_PATH.write_text(json.dumps({'updated_at': now, 'items': items}, ensure_ascii=False, indent=2), encoding='utf-8')
    STATUS_PATH.write_text(json.dumps({'updated_at': now, 'mode': mode, 'sources': statuses}, ensure_ascii=False, indent=2), encoding='utf-8')
    CHAINS_PATH.write_text(json.dumps({'updated_at': now, 'chains': build_chains(items)}, ensure_ascii=False, indent=2), encoding='utf-8')
    print('mode=', mode, 'items=', len(items), 'sources=', len(statuses),
          'failed=', sum(1 for x in statuses if x['status'] == '抓取失败'),
          'partial=', sum(1 for x in statuses if x['status'] == '部分失败'),
          'new=', sum(x['new_items'] for x in statuses),
          'changed=', sum(x['changed_items'] for x in statuses),
          'refreshed=', sum(x['refreshed_items'] for x in statuses))


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--mode', choices=['quick', 'deep'], default='quick')
    args = ap.parse_args()
    main(args.mode)
